import os, csv, sys, time, shutil, subprocess, threading, Queue, datetime
from time import gmtime, strftime
from openpyxl import Workbook
import xlsxwriter


wb = Workbook()
ws = wb.active
workbook = xlsxwriter.Workbook('chart_line.xlsx')
worksheet = workbook.add_worksheet()

with open('battery_test.csv') as csvfile:
    testResults = csv.DictReader(csvfile)
    currentList = []

    for row in testResults:
        current = row['Main Avg Power (mW)']
        currentList.append(int(current))


data = currentList
worksheet.write_column('A1', data)
data1 = currentList.reverse()
worksheet.write_column('B1', data)


# Create a new chart object.
chart = workbook.add_chart({'type': 'line'})

# Add a series to the chart.
chart.add_series({'name': 'yabro', 'values': '=Sheet1!$A$1:$A$'+str(len(currentList)), 'line': {'width': 1.25}})
chart.add_series({'name': 'chrome', 'values': '=Sheet1!$B$1:$B$'+str(len(currentList)), 'line': {'width': 1.25}})

chart.set_size({'width': 1800, 'height': 800})

# Insert the chart into the worksheet.
worksheet.insert_chart('A1', chart)

workbook.close()